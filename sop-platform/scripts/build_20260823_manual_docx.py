from __future__ import annotations

import os
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(
    os.getenv(
        "SOP_MANUAL_SOURCE",
        ROOT / "docs" / "27_宁波视觉SOP_CVAT闭环平台完整说明书_20260823.md",
    )
)
OUTPUT = SOURCE.with_suffix(".docx")


def set_run_font(run, size: float = 12, bold: bool | None = None) -> None:
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "KaiTi")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def configure_style(style, size: float, bold: bool, before: float, after: float) -> None:
    style.font.name = "Times New Roman"
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "KaiTi")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.keep_with_next = True


def add_inline(paragraph, text: str, size: float = 12) -> None:
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        marked = (part.startswith("`") and part.endswith("`")) or (part.startswith("**") and part.endswith("**"))
        content = part[1:-1] if part.startswith("`") else part[2:-2] if part.startswith("**") else part
        set_run_font(paragraph.add_run(content), size, marked)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "在 Word 中右键更新域以刷新目录"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, text, separate, placeholder, end])
    set_run_font(run, 11)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE")


def build() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ").strip()
    version = next((line.removeprefix("版本：") for line in lines if line.startswith("版本：")), "2026-08-23")
    sop_address = next(
        (line.removeprefix("SOP 平台：").replace("`", "") for line in lines if line.startswith("SOP 平台：")),
        "SOP 平台",
    )
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.4)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "KaiTi")
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(0.74)
    configure_style(document.styles["Heading 1"], 18, True, 16, 8)
    configure_style(document.styles["Heading 2"], 15, True, 12, 6)
    configure_style(document.styles["Heading 3"], 13, True, 10, 4)

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Cm(5.0)
    cover.paragraph_format.space_after = Cm(2.0)
    set_run_font(cover.add_run(title), 24, True)
    for value in (version, f"{sop_address} · CVAT 8081 · MediaMTX 8554", "宁波视觉 SOP 项目组"):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(12)
        set_run_font(paragraph.add_run(value), 14, value == version)
    document.add_page_break()

    toc_title = document.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.space_after = Pt(16)
    set_run_font(toc_title.add_run("目录"), 20, True)
    toc = document.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u')
    document.add_page_break()

    in_code = False
    started = False
    for raw in lines[1:]:
        line = raw.rstrip()
        if line.startswith("## "):
            started = True
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:], 18)
        elif line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:], 15)
        elif line.startswith("#### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_inline(paragraph, line[5:], 13)
        elif not started:
            continue
        elif line.startswith("```"):
            in_code = not in_code
            continue
        elif not line:
            continue
        elif re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, line[2:])
        elif line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, line[2:])
        else:
            paragraph = document.add_paragraph()
            if in_code:
                paragraph.paragraph_format.left_indent = Cm(0.74)
                paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, line)
        if paragraph.style.name.startswith("Heading"):
            continue
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.space_after = Pt(6)

    for current in document.sections:
        add_page_number(current)
    document.core_properties.title = title
    document.core_properties.subject = "宁波视觉 SOP、CVAT、AI 标注、模型训练和 MES 闭环操作说明"
    document.core_properties.author = "宁波视觉 SOP 项目组"
    document.core_properties.keywords = "SOP,CVAT,AI标注,YOLO,MediaMTX,MES"
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build())
