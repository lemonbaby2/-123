#!/usr/bin/env python3
"""Render the meeting Markdown manual to a screenshot-rich Word document."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_font(run, size: float = 12, bold: bool | None = None, color: str | None = None) -> None:
    run.font.name = "Times New Roman"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), "楷体")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_font(style, size: float, bold: bool, before: float, after: float) -> None:
    style.font.name = "Times New Roman"
    fonts = style._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:eastAsia"), "楷体")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = 1.25
    style.paragraph_format.keep_with_next = True


def add_inline(paragraph, text: str, size: float = 12) -> None:
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text):
        if not part:
            continue
        code = part.startswith("`") and part.endswith("`")
        strong = part.startswith("**") and part.endswith("**")
        content = part[1:-1] if code else part[2:-2] if strong else part
        run = paragraph.add_run(content)
        set_font(run, size, code or strong)
        if code:
            run.font.name = "Consolas"
            run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "等线")


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    code = OxmlElement("w:instrText")
    code.set(qn("xml:space"), "preserve")
    code.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, code, separate, value, end])
    set_font(run, 10)


def add_page_number(section) -> None:
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(lines) and lines[index].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def build(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    title = lines[0].removeprefix("# ").strip()
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    normal = document.styles["Normal"]
    style_font(normal, 11.5, False, 0, 6)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)
    style_font(document.styles["Heading 1"], 18, True, 16, 8)
    style_font(document.styles["Heading 2"], 15, True, 12, 6)
    style_font(document.styles["Heading 3"], 13, True, 10, 4)
    style_font(document.styles["Caption"], 10, False, 4, 8)

    cover = document.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Cm(4.8)
    set_font(cover.add_run(title), 25, True, "0A2847")
    for value in ("2026-08-24 培训版", "SOP 8116 · CVAT 8081", "宁波视觉 SOP 项目组"):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_before = Pt(10)
        set_font(p.add_run(value), 14, value.endswith("培训版"), "1F8168")
    document.add_page_break()

    toc_title = document.add_paragraph()
    toc_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(toc_title.add_run("目录"), 20, True, "0A2847")
    toc = document.add_paragraph()
    toc.paragraph_format.first_line_indent = Cm(0)
    add_field(toc, 'TOC \\o "1-3" \\h \\z \\u', "打开 Word 后右键目录并选择“更新域”")
    document.add_page_break()

    started = False
    index = 1
    while index < len(lines):
        line = lines[index].rstrip()
        if line.startswith("## "):
            if started and re.match(r"## (4|7|8|9|10|11|12|13|14|15|16)\.", line):
                document.add_page_break()
            started = True
            paragraph = document.add_paragraph(style="Heading 1")
            add_inline(paragraph, line[3:], 18)
        elif line.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            add_inline(paragraph, line[4:], 15)
        elif line.startswith("#### "):
            paragraph = document.add_paragraph(style="Heading 3")
            add_inline(paragraph, line[5:], 13)
        elif not started or not line:
            index += 1
            continue
        elif re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", line):
            match = re.fullmatch(r"!\[([^]]+)]\(([^)]+)\)", line)
            assert match
            caption, relative = match.groups()
            image_path = (source.parent / relative).resolve()
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.first_line_indent = Cm(0)
            picture.add_run().add_picture(str(image_path), width=Cm(16.1))
            cap = document.add_paragraph(style="Caption")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.first_line_indent = Cm(0)
            set_font(cap.add_run(caption), 10)
        elif line.startswith("|") and index + 1 < len(lines) and lines[index + 1].strip().startswith("|"):
            rows, next_index = parse_table(lines, index)
            if rows:
                columns = max(len(row) for row in rows)
                table = document.add_table(rows=len(rows), cols=columns)
                table.style = "Table Grid"
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                set_repeat_table_header(table.rows[0])
                for row_no, row in enumerate(rows):
                    for col_no in range(columns):
                        cell = table.cell(row_no, col_no)
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        cell.text = ""
                        paragraph = cell.paragraphs[0]
                        paragraph.paragraph_format.first_line_indent = Cm(0)
                        add_inline(paragraph, row[col_no] if col_no < len(row) else "", 9.5)
                        if row_no == 0:
                            for run in paragraph.runs:
                                run.bold = True
                index = next_index
                continue
        elif re.match(r"^\d+\. ", line):
            paragraph = document.add_paragraph(style="List Number")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, re.sub(r"^\d+\. ", "", line))
        elif line.startswith("- "):
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.first_line_indent = Cm(0)
            add_inline(paragraph, line[2:])
        elif line.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.74)
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.keep_together = True
            add_inline(paragraph, line[2:])
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(130, 45, 45)
        else:
            paragraph = document.add_paragraph()
            add_inline(paragraph, line)

        if "paragraph" in locals() and not paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(6)
        index += 1

    for current in document.sections:
        add_page_number(current)
    settings = document.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    document.core_properties.title = title
    document.core_properties.subject = "明日会议工人标注、训练、推理与实时 SOP 培训"
    document.core_properties.author = "宁波视觉 SOP 项目组"
    document.core_properties.keywords = "SOP,CVAT,标注,Track,Reviewer,训练,MES"
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.source, args.output)
    print(args.output)


if __name__ == "__main__":
    main()
